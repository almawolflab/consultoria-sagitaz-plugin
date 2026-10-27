#!/usr/bin/env python3
"""
plantilla_io.py — I/O de la plantilla estándar de CRM (Excel) para la parametrización (AW-161).

Opera sobre la **plantilla base con estilo** (`assets/plantilla_crm_base.xlsx`, derivada de la estándar de
SagitaZ, sin datos de cliente): la carga, inyecta en ella los campos propuestos (en el bloque «Campos
Personalizados», en amarillo + `Origen=DERCAS`) y la relee. Hereda el estilo del base (colores, anchos,
pestañas, cabeceras agrupadas). El descriptor `assets/plantilla_crm.json` aporta el mapeo hoja→variante y
el vocabulario de tipos.

Depende de **openpyxl** (Python puro, vendorizado). Ver contrato `docs/27_contrato_plantilla_parametrizacion.md`.
NOTA: el logo (imagen embebida) requiere Pillow para conservarse al re-guardar con openpyxl; en Catalyst
(sin Pillow) se pierde en la pre-marcada → re-inyección por zip pendiente (mejora posterior). `generar_blank`
copia el fichero byte a byte, así que el blank sí conserva el logo.
"""
import json
import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

BASE = Path(__file__).resolve().parent.parent          # plugin-claude/scripts/
ASSETS = BASE / "assets"

ORIGEN_DERCAS = "DERCAS"
_AMARILLO = "FFFF00"
_FILL_AMARILLO = PatternFill(fill_type="solid", fgColor=_AMARILLO)
_YELLOWS = {"FFFF00", "FFF2CC", "FFE699", "FFD966", "FFEB9C", "FFFFCC", "FFFFFF00"}
_MAX_COL = 13   # cap de columnas (la base trae max_column=1050 en variante A por artefacto de formato)


# ---------------------------------------------------------------- descriptor + tipos

def cargar_descriptor(path=None):
    p = Path(path) if path else (ASSETS / "plantilla_crm.json")
    return json.loads(p.read_text(encoding="utf-8"))


def _base_path(descriptor):
    return ASSETS / (descriptor.get("meta", {}).get("base_xlsx") or "plantilla_crm_base.xlsx")


def _norm_tipo(tipo):
    return " ".join(str(tipo or "").strip().lower().split())


def clasificar_tipo(tipo, descriptor):
    """Etiqueta de tipo de la plantilla → (zoho_data_type, bucket). Desconocido → (None, 'manual')."""
    info = descriptor.get("tipos", {}).get(_norm_tipo(tipo))
    if info:
        return info.get("zoho"), info.get("bucket")
    return None, "manual"


ZOHO_DISPLAY = {
    "text": "Texto", "textarea": "Texto amplio", "picklist": "Lista de selección",
    "multiselectpicklist": "Lista de selección múltiple", "integer": "Número",
    "bigint": "Número entero largo", "double": "Decimal", "currency": "Moneda",
    "date": "Fecha", "datetime": "Fecha/hora", "boolean": "Casilla de verificación",
    "email": "Correo electrónico", "phone": "Teléfono", "website": "URL",
    "ownerlookup": "Usuario", "lookup": "Lookup", "autonumber": "Autonumérico",
    "formula": "Fórmula", "imageupload": "Carga de imagen",
}


def display_de_zoho(zoho, descriptor=None):
    """`data_type` de Zoho → etiqueta ES para la plantilla (p.ej. picklist → 'Lista de selección')."""
    return ZOHO_DISPLAY.get(zoho, zoho or "Texto")


def bucket_de_zoho(zoho, descriptor):
    """`data_type` de Zoho → bucket (v1|mas_adelante|manual), consultando el vocabulario `tipos`."""
    for info in descriptor.get("tipos", {}).values():
        if info.get("zoho") == zoho:
            return info.get("bucket", "manual")
    return "manual"


# ---------------------------------------------------------------- layout de la base

def _col_por_variante(variante):
    """Posición (1-based) de las columnas relevantes según la variante de la hoja."""
    if variante == "A":
        return {"modulo": 1, "seccion": 3, "campo": 4, "tipo": 5, "valores": 8, "origen": 13}
    return {"modulo": 1, "seccion": None, "campo": 2, "tipo": 3, "valores": 6, "origen": 13}


def _n(v):
    return str(v).strip().lower() if v is not None else ""


def _v(ws, r, c):
    return str(ws.cell(r, c).value).strip() if (c and ws.cell(r, c).value is not None) else ""


def _headers_y_marcador(ws):
    """Localiza (fila header estándar, fila marcador 'Campos Personalizados', fila header custom)."""
    std_h = cust_h = marker = None
    for r in range(1, min(ws.max_row, 80) + 1):
        row_txt = " ".join(_n(ws.cell(r, c).value) for c in range(1, 6))
        if "módulo" in row_txt and "campo" in row_txt:
            if std_h is None:
                std_h = r
            elif cust_h is None:
                cust_h = r
        if marker is None and ("personalizado" in row_txt or "campos nuevos" in row_txt):
            marker = r
    return std_h, marker, cust_h


def _es_amarillo(cell):
    f = cell.fill
    if not f or f.patternType != "solid":
        return False
    rgb = getattr(f.fgColor, "rgb", None)
    return isinstance(rgb, str) and rgb[-6:].upper() in _YELLOWS


# ---------------------------------------------------------------- generar / pre-marcar

def generar_blank(descriptor, out_xlsx):
    """Plantilla en blanco = copia byte a byte del base (conserva logo y estilo)."""
    shutil.copy(str(_base_path(descriptor)), str(out_xlsx))
    return out_xlsx


def _set(ws, r, c, val, ref=None):
    cell = ws.cell(r, c, val)
    if ref is not None:
        cell._style = copy(ws.cell(ref[0], ref[1])._style)   # heredar el estilo de una fila estándar
    return cell


def escribir_premarcada(descriptor, propuestas, out_xlsx):
    """Carga la base e inyecta las `propuestas` en el bloque 'Campos Personalizados' de cada módulo,
    en amarillo + Origen=DERCAS. `propuestas`: {api_name: [{campo,tipo,seccion?,valores?}]}."""
    wb = load_workbook(str(_base_path(descriptor)))
    for modulo in descriptor["modulos"]:
        hoja = modulo["hoja"]
        if hoja not in wb.sheetnames:
            continue
        rows = (propuestas or {}).get(modulo["api_name"]) or []
        ws = wb[hoja]
        cols = _col_por_variante(modulo["variante"])
        std_h, _marker, cust_h = _headers_y_marcador(ws)
        if not cust_h:
            continue
        ws.cell(cust_h, cols["origen"], "Origen")            # cabecera Origen
        ref = (std_h or 1) + 1                                 # fila estándar de referencia (para el estilo)
        modulo_val = _v(ws, ref, cols["modulo"]) or hoja
        rr = cust_h + 1
        for p in rows:
            _set(ws, rr, cols["modulo"], modulo_val, ref=(ref, cols["modulo"]))
            if cols["seccion"] and p.get("seccion"):
                _set(ws, rr, cols["seccion"], p["seccion"], ref=(ref, cols["seccion"]))
            _set(ws, rr, cols["campo"], p.get("campo", ""), ref=(ref, cols["campo"]))
            _set(ws, rr, cols["tipo"], p.get("tipo", ""), ref=(ref, cols["tipo"]))
            if p.get("valores"):
                _set(ws, rr, cols["valores"], "\n".join(p["valores"]), ref=(ref, cols["valores"]))
            _set(ws, rr, cols["origen"], ORIGEN_DERCAS)
            ws.cell(rr, cols["campo"]).fill = _FILL_AMARILLO
            rr += 1
    wb.save(str(out_xlsx))
    return out_xlsx


# ---------------------------------------------------------------- leer plantilla rellena

def leer_plantilla(xlsx_path, descriptor):
    """Lee una plantilla (base o rellena) y devuelve su contenido estructurado:
       {modulos: {api_name: {hoja, estandar:[{campo,nota}], nuevos:[{campo,tipo,seccion,valores,origen,amarillo}]}}}."""
    wb = load_workbook(str(xlsx_path), data_only=True)
    by_hoja = {m["hoja"]: m for m in descriptor["modulos"]}
    out = {"modulos": {}}
    for hoja, modulo in by_hoja.items():
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        cols = _col_por_variante(modulo["variante"])
        std_h, marker, cust_h = _headers_y_marcador(ws)
        estandar, nuevos = [], []
        if std_h:
            fin = marker or cust_h or (ws.max_row + 1)
            for r in range(std_h + 1, fin):
                campo = _v(ws, r, cols["campo"])
                if not campo or campo.lower() == "campo":
                    continue
                estandar.append({"campo": campo, "nota": _v(ws, r, cols["valores"])})
        if cust_h:
            for r in range(cust_h + 1, ws.max_row + 1):
                campo = _v(ws, r, cols["campo"])
                if not campo or campo.lower() == "campo":
                    continue
                valores = _v(ws, r, cols["valores"])
                nuevos.append({
                    "campo": campo,
                    "tipo": _v(ws, r, cols["tipo"]),
                    "seccion": _v(ws, r, cols["seccion"]) if cols["seccion"] else "",
                    "valores": [x.strip() for x in valores.split("\n") if x.strip()],
                    "origen": _v(ws, r, cols["origen"]),
                    "amarillo": _es_amarillo(ws.cell(r, cols["campo"])),
                })
        out["modulos"][modulo["api_name"]] = {"hoja": hoja, "estandar": estandar, "nuevos": nuevos}
    return out


if __name__ == "__main__":
    import sys
    d = cargar_descriptor()
    out = sys.argv[1] if len(sys.argv) > 1 else "plantilla_crm_blank.xlsx"
    generar_blank(d, out)
    print("Plantilla en blanco (copia del base con estilo):", out)
