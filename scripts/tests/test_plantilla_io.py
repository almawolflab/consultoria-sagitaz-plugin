"""Tests de la capa de I/O de la plantilla CRM (plugin-claude/scripts/parametrizacion/)."""
import pathlib

import pytest

SCRIPTS = pathlib.Path(__file__).resolve().parents[1]  # plugin-claude/scripts/

pytest.importorskip("openpyxl")

from parametrizacion import plantilla_io as P


def test_descriptor_valido():
    d = P.cargar_descriptor()
    assert len(d["modulos"]) == 15
    for m in d["modulos"]:
        assert m["api_name"] and m["variante"] in ("A", "B") and m["hoja"]
    assert d["tipos"] and d["planes"]["estandar"] == 10
    assert "Leads" in {m["api_name"] for m in d["modulos"]}
    # el asset base existe
    assert (SCRIPTS / "assets/plantilla_crm_base.xlsx").exists()


def test_clasificar_tipo():
    d = P.cargar_descriptor()
    assert P.clasificar_tipo("Texto", d) == ("text", "v1")
    assert P.clasificar_tipo("lista de selección", d) == ("picklist", "v1")
    assert P.clasificar_tipo("Campo de tipo fórmula", d)[1] == "mas_adelante"
    assert P.clasificar_tipo("tipo inexistente", d) == (None, "manual")


def test_generar_blank_es_copia_del_base(tmp_path):
    d = P.cargar_descriptor()
    out = tmp_path / "blank.xlsx"
    P.generar_blank(d, str(out))
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    for hoja in ("Posible clientes", "Oportunidades", "Tareas", "Notificaciones", "Workflow", "Instrucciones"):
        assert hoja in wb.sheetnames
    # las 2 hojas custom se quitaron del estándar
    assert "Gestión de Cobro" not in wb.sheetnames
    r = P.leer_plantilla(str(out), d)
    assert len(r["modulos"]["Leads"]["estandar"]) >= 25   # campos estándar del base
    assert r["modulos"]["Leads"]["nuevos"] == []           # blank: sin campos nuevos


def test_roundtrip_premarcada(tmp_path):
    d = P.cargar_descriptor()
    prop = {
        "Leads": [
            {"campo": "JL_Origen", "tipo": "Lista de selección",
             "seccion": "Información del contacto", "valores": ["Web", "Referido"]},
            {"campo": "JL_DNI", "tipo": "Texto"},
        ],
        "Deals": [{"campo": "JL_Estado_proc", "tipo": "Lista de selección", "valores": ["1a", "2a"]}],
    }
    out = tmp_path / "pre.xlsx"
    P.escribir_premarcada(d, prop, str(out))
    r = P.leer_plantilla(str(out), d)
    leads = r["modulos"]["Leads"]["nuevos"]
    assert len(leads) == 2
    n0 = next(n for n in leads if n["campo"] == "JL_Origen")
    assert n0["tipo"] == "Lista de selección"
    assert n0["seccion"] == "Información del contacto"
    assert n0["valores"] == ["Web", "Referido"]
    assert n0["origen"] == "DERCAS" and n0["amarillo"] is True
    assert len(r["modulos"]["Deals"]["nuevos"]) == 1
    # los campos estándar del base se conservan
    assert len(r["modulos"]["Leads"]["estandar"]) >= 25


def test_amarillo_solo_en_nuevos(tmp_path):
    """El base no tiene amarillos (reset); solo las filas propuestas quedan marcadas."""
    d = P.cargar_descriptor()
    out = tmp_path / "b.xlsx"
    P.escribir_premarcada(d, {"Leads": [{"campo": "JL_X", "tipo": "Texto"}]}, str(out))
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    ws = wb["Posible clientes"]
    amarillas = sum(1 for row in ws.iter_rows() for c in row if P._es_amarillo(c))
    assert amarillas == 1
